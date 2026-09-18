"""Таблица Excel для ручной разметки (калибровка на реальных фото). Инструмент разработчика.

Для каждой папки — лист: фото в порядке съёмки, ссылка для открытия фото, предложение
программы и колонка «Верная категория» с выпадающим списком категорий эталона.
Пользователь исправляет неверные строки, а tools/evaluate_labels.py считает по таблице точность.

    python tools/make_labels.py --reference <эталон> <папка> [<папка> ...] --out logs/разметка/файл.xlsx

Фото только читаются. Таблица содержит имена файлов и категории, изображений в ней нет.
"""

import argparse
import os
import sys
from urllib.parse import quote

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from photo_renamer import features as feat  # noqa: E402
from photo_renamer import matcher  # noqa: E402
from photo_renamer.reference import load_reference  # noqa: E402
from photo_renamer.scanner import list_photos  # noqa: E402

NOT_A_CATEGORY = "Не относится ни к одной категории"
HEADER_ROW = 4
COLUMNS = ["№", "Файл (щелчок — открыть фото)", "Предложение программы", "Уверенность, %", "Статус",
           "Верная категория (проверьте!)", "Комментарий"]
TRUTH_COL = 6
STATUS_TEXT = {
    matcher.STATUS_OK: "уверенно",
    matcher.STATUS_SIDE_RULE: "по правилу сторон",
    matcher.STATUS_LOW_CONFIDENCE: "НЕ УВЕРЕНА",
    matcher.STATUS_NO_MATCH: "НЕ ПОХОЖЕ НИ НА ЧТО",
    matcher.STATUS_UNREADABLE: "не прочитано",
    matcher.STATUS_ALREADY_NAMED: "уже названо",
}
DOUBT = PatternFill("solid", fgColor="FFF2CC")

INSTRUCTION = """Как заполнить таблицу

1. На листе каждой папки строки идут в порядке съёмки (по имени файла).
2. Щелчок по имени файла открывает фото.
3. Колонка «Верная категория» уже заполнена предложением программы.
   Проверьте КАЖДУЮ строку. Если категория неверная, выберите верную из выпадающего списка.
4. Жёлтым выделены строки, где программа не уверена. Их проверьте особенно внимательно.
5. Если фото не относится ни к одной категории эталона (посторонний снимок, брак),
   выберите «{none}».
6. Номера экземпляров и сторон указывать не нужно, только категорию.
7. В «Комментарий» можно написать что угодно (например, «размыто», «вторая сторона»).
8. Не меняйте имена файлов, порядок строк и названия листов. Сохраните файл в формате .xlsx.
""".format(none=NOT_A_CATEGORY)


def file_link(path):
    return "file://" + quote(os.path.abspath(path))


def build(reference_dir, folders, out_path, white_balance):
    reference = load_reference(reference_dir)
    extractor = feat.FeatureExtractor(white_balance=white_balance)
    print("Эталон: %d фото" % len(reference.photos))
    model = matcher.ReferenceModel.build(reference, extractor)
    categories = [c.name for c in reference.categories.values()] + [NOT_A_CATEGORY]

    wb = Workbook()
    info = wb.active
    info.title = "Инструкция"
    for i, line in enumerate(INSTRUCTION.splitlines(), start=1):
        info.cell(row=i, column=1, value=line)
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 110
    base = len(INSTRUCTION.splitlines()) + 2
    info.cell(row=base, column=1, value="Эталон: %s" % os.path.abspath(reference_dir))
    info.cell(row=base + 1, column=1, value="Выравнивание цвета по фону при подготовке предложений: %s"
              % ("да" if white_balance else "нет"))

    cat_sheet = wb.create_sheet("Категории")
    for i, name in enumerate(categories, start=1):
        cat_sheet.cell(row=i, column=1, value=name)
    cat_sheet.column_dimensions["A"].width = 45
    list_ref = "'Категории'!$A$1:$A$%d" % len(categories)

    for folder in folders:
        names = list_photos(folder).photos
        print("Папка %s: %d фото" % (folder, len(names)))
        result = matcher.match_folder(model, folder, names)
        ws = wb.create_sheet("Папка %s" % os.path.basename(os.path.normpath(folder))[:24])
        ws["A1"] = "Папка:"
        ws["B1"] = os.path.abspath(folder)
        ws["A2"] = "Проверьте колонку F в каждой строке; жёлтые строки — программа не уверена."
        ws["A2"].font = Font(italic=True)
        for col, title in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=HEADER_ROW, column=col, value=title)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        validation = DataValidation(type="list", formula1="=" + list_ref, allow_blank=False,
                                    showErrorMessage=True, errorTitle="Неизвестная категория",
                                    error="Выберите категорию из списка")
        ws.add_data_validation(validation)

        for k, p in enumerate(result.photos, start=1):
            row = HEADER_ROW + k
            if p.identified:
                suggestion = p.category.name
            elif p.candidates:
                suggestion = p.candidates[0].category
            else:
                suggestion = ""
            ws.cell(row=row, column=1, value=k)
            link = ws.cell(row=row, column=2, value=p.filename)
            link.hyperlink = file_link(os.path.join(folder, p.filename))
            link.font = Font(color="0563C1", underline="single")
            ws.cell(row=row, column=3, value=suggestion)
            ws.cell(row=row, column=4, value=round(p.confidence, 1))
            ws.cell(row=row, column=5, value=STATUS_TEXT.get(p.status, p.status))
            truth = ws.cell(row=row, column=TRUTH_COL, value=suggestion or NOT_A_CATEGORY)
            validation.add(truth)
            if p.status not in (matcher.STATUS_OK,):
                for col in range(1, len(COLUMNS) + 1):
                    ws.cell(row=row, column=col).fill = DOUBT
        for col, width in zip("ABCDEFG", (5, 26, 34, 13, 22, 38, 30)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    print("Сохранено: %s" % out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--reference", "-r", required=True)
    parser.add_argument("folders", nargs="+")
    parser.add_argument("--out", required=True)
    parser.add_argument("--white-balance", action="store_true", help="выравнивать цвета по фону")
    args = parser.parse_args()
    build(args.reference, args.folders, args.out, args.white_balance)


if __name__ == "__main__":
    main()
